from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "OCTO_EARTH_PITCH_DECK.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = "0B1725"
PANEL = "12283B"
PANEL2 = "18354B"
WHITE = "F2F6F8"
MUTED = "ADC0CE"
TEAL = "53D2D0"
GOLD = "F0B765"
BLUE = "86BDF8"
RED = "FF9C83"
GREEN = "8DE0B2"
LINE = "2A465B"


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def set_bg(slide):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = rgb(NAVY)


def textbox(slide, x, y, w, h, text, size=16, color=WHITE, bold=False,
            font="Arial", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            margin=0.04, italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run(); run.text = text
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return shape


def rect(slide, x, y, w, h, fill=PANEL, line=LINE, radius=True, line_width=1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line); shape.line.width = Pt(line_width)
    return shape


def oval(slide, x, y, w, h, fill=None, line=TEAL, line_width=1.5):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    sh.line.color.rgb = rgb(line); sh.line.width = Pt(line_width)
    return sh


def line(slide, x1, y1, x2, y2, color=TEAL, width=2.0, dash=None):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = rgb(color); sh.line.width = Pt(width)
    if dash:
        sh.line.dash_style = dash
    return sh


def title(slide, idx, section, heading, subtitle, timing):
    textbox(slide, .62, .23, 7.6, .24, f"{idx:02d}  /  {section.upper()}", 9, TEAL, True)
    textbox(slide, 10.32, .23, 2.35, .24, timing, 9, MUTED, False, align=PP_ALIGN.RIGHT)
    textbox(slide, .62, .56, 12.05, .52, heading, 26, WHITE, True)
    if subtitle:
        textbox(slide, .65, 1.10, 12.0, .34, subtitle, 12, MUTED)
    line(slide, .65, 1.52, 12.68, 1.52, LINE, 1.0)
    line(slide, .65, 7.08, 12.68, 7.08, LINE, 0.9)
    textbox(slide, .65, 7.13, 7.8, .18, "OCTO-EARTH  |  hypothetical model · not a feasibility claim", 8, MUTED)
    textbox(slide, 12.0, 7.13, .65, .18, f"{idx:02d}", 8, MUTED, align=PP_ALIGN.RIGHT)


def bullet_lines(slide, x, y, w, h, items, size=13, color=WHITE, gap=5, bullet_color=TEAL):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.05); tf.margin_right = Inches(.02); tf.margin_top = Inches(.03); tf.margin_bottom = Inches(.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.name = "Arial"; p.font.size = Pt(size); p.font.color.rgb = rgb(color)
        p.space_after = Pt(gap); p.level = 0
    return shape


def label(slide, x, y, w, text, color=TEAL):
    textbox(slide, x, y, w, .24, text.upper(), 9, color, True)


def card_text(slide, x, y, w, h, heading, body, accent=TEAL, hsize=14, bsize=11):
    rect(slide, x, y, w, h)
    textbox(slide, x+.18, y+.14, w-.36, .27, heading, hsize, accent, True)
    textbox(slide, x+.18, y+.48, w-.36, h-.58, body, bsize, WHITE)


# 1 — cover / question
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
title(s, 1, "The question", "Can we move the orbit without breaking the physics?", "OCTO-EARTH is a physics-led simulation prototype—not a claim of planetary feasibility.", "0:00–0:35")
textbox(s, .82, 1.88, 6.15, .55, "1 AU  →  2.7 AU", 31, GOLD, True)
textbox(s, .86, 2.52, 5.9, .64, "Earth–Moon system · ≤10,000 simulated years\nwith a simplified climate objective", 17, WHITE)
rect(s, .82, 3.56, 5.85, 1.14, PANEL2, LINE)
textbox(s, 1.02, 3.72, 5.45, .8, "Earth changes orbit only by reaction to finite ejected mass. The agent never edits velocity directly and never creates energy.", 14, WHITE, True)
textbox(s, .86, 5.05, 5.8, .65, "p40: 48 tests pass · 1 Torch-dependent device test blocked\nPhysics check blocked (Stable-Baselines3 / Torch missing)\nNo current PPO evaluation · no hardware certification", 10, MUTED)
# schematic solar system at right
cx, cy = 10.00, 4.13
oval(s, cx-2.43, cy-1.60, 4.86, 3.20, line=GOLD, line_width=1.4)
oval(s, cx-1.20, cy-.80, 2.40, 1.60, line=TEAL, line_width=1.6)
oval(s, cx-.49, cy-.32, .98, .64, fill=GOLD, line=GOLD)
oval(s, cx-.12, cy-.12, .24, .24, fill="FFF4C9", line="FFF4C9", line_width=.5)
oval(s, cx-1.18, cy-.08, .16, .16, fill=BLUE, line=BLUE, line_width=.5)
oval(s, cx-1.08, cy+.04, .045, .045, fill=WHITE, line=WHITE, line_width=.4)
oval(s, cx+2.32, cy-.08, .16, .16, fill=BLUE, line=BLUE, line_width=.5)
oval(s, cx+2.42, cy+.04, .045, .045, fill=WHITE, line=WHITE, line_width=.4)
textbox(s, 7.7, 1.82, 1.9, .26, "Sun", 10, GOLD, True)
textbox(s, 8.05, 2.62, 2.7, .3, "1 AU · reference", 10, TEAL, True)
textbox(s, 10.45, 2.28, 2.1, .32, "2.7 AU · target", 10, GOLD, True)
textbox(s, 9.15, 5.95, 2.0, .24, "SCHEMATIC · not to scale", 9, MUTED, True, align=PP_ALIGN.CENTER)

# 2 — physics / reaction mass
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
title(s, 2, "Momentum", "The only push is a real reaction", "Mass is finite; fragments retain their velocities; all eight arms throw one packet together per event.", "0:35–1:25")
# diagram left
rect(s, .78, 1.83, 5.95, 4.85)
# Earth
oval(s, 2.85, 3.48, 1.22, 1.22, fill="3077A8", line="8CC7F2", line_width=1.4)
textbox(s, 2.96, 3.88, 1.0, .28, "EARTH", 11, WHITE, True, align=PP_ALIGN.CENTER)
# ring / arms
oval(s, 2.50, 3.13, 1.93, 1.93, line=TEAL, line_width=1.3)
root_pts = [(3.47,3.12),(4.24,3.50),(4.42,4.10),(4.05,4.78),(3.47,5.06),(2.72,4.72),(2.52,4.10),(2.83,3.50)]
for x,y in root_pts:
    line(s, 3.47, 4.09, x, y, color=TEAL, width=1.0)
# asteroid and ejected packet
oval(s, 3.35, 2.10, .24, .24, fill=GOLD, line=GOLD, line_width=.5)
textbox(s, 2.42, 2.34, 2.1, .3, "finite asteroid", 10, GOLD, True, align=PP_ALIGN.CENTER)
line(s, 3.48, 2.02, 4.72, 1.70, color=GOLD, width=2.5)
textbox(s, 4.70, 1.52, 1.4, .3, "packet →", 11, GOLD, True)
line(s, 2.72, 4.18, 1.65, 4.18, color=BLUE, width=2.4)
textbox(s, 1.03, 3.84, 1.5, .46, "Earth–Moon\nrecoil", 11, BLUE, True, align=PP_ALIGN.CENTER)
textbox(s, 1.08, 5.70, 5.1, .62, "2 R_E arms · one joint packet event\nfragments / impacts / escapes remain in the ledger", 10, MUTED, align=PP_ALIGN.CENTER)
# explanation right
card_text(s, 7.05, 1.90, 5.45, 1.13, "1 · Finite reservoir", "Reaction mass is debited packet by packet; the reservoir cannot go negative.", GOLD, 13, 11)
card_text(s, 7.05, 3.20, 5.45, 1.13, "2 · Equal and opposite", "The Earth–Moon system recoils under momentum conservation. No teleport and no direct velocity command.", BLUE, 13, 11)
card_text(s, 7.05, 4.50, 5.45, 1.13, "3 · Energy stays explicit", "Launcher draw, packet kinetic energy and heat are tracked. No free energy; reward is not a physical metric.", TEAL, 13, 11)
textbox(s, 7.16, 5.92, 5.15, .55, "Moon gravity is present in the fast and high-fidelity models. High-fidelity runs track explicit fragments.", 10, MUTED)

# 3 — three individual enablers
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
title(s, 3, "Enablers", "Three single assumptions. No bundled systems.", "Selected for mission criticality; support hardware and gaps are stated separately.", "1:25–2:10")
card_text(s, .82, 1.90, 3.86, 2.42, "E1 · FINITE ASTEROID ACCESS", "Access to a configured 3×10²² kg reaction-mass reservoir.\n\nWhy: without reaction mass, modeled recoil cannot transfer momentum to Earth.\n\nNo capture / docking is simulated.", GOLD, 13, 11)
card_text(s, 4.76, 1.90, 3.86, 2.42, "E2 · THIN-FILM PV FIELD", "Five million collectors at 0.5 AU, intercepting configured 30% of solar output.\n\nWhy: supplies the modeled electrical source.\n\nScale and construction are unverified.", TEAL, 13, 11)
card_text(s, 8.70, 1.90, 3.86, 2.42, "E3 · SPACE-MIRROR RING", "Reflects sunlight into the simplified climate model at 2.7 AU.\n\nWhy: direct sunlight falls; mirrors redirect rather than create energy.\n\nNo real formation-flight design.", BLUE, 13, 11)
rect(s, .82, 4.66, 11.74, 1.54, PANEL2, LINE)
label(s, 1.02, 4.82, 5.4, "Separate supporting systems — not additional enablers", GOLD)
textbox(s, 1.03, 5.18, 11.22, .78, "Launcher ring + 8 arms · power conditioning / reserve · directional radiators · microwave arrays + rectennas · collector station-keeping + sunlight lanes · factories / Mercury feedstock · mirror orbit / formation / pointing.\nEach item has its own performance gap; “supporting” does not mean ordinary, feasible, safe or certified.", 11, WHITE)
textbox(s, .92, 6.42, 11.4, .3, "Selected enablers: asteroid access  |  photovoltaic generation  |  mirror reflection", 10, MUTED, italic=True)

# 4 — quantitative scale
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
title(s, 4, "Scale", "The hypothetical preset is planetary-scale", "Every value is a model estimate from one configuration—not an engineered capability.", "2:10–3:20")
metrics = [
    (.82,1.86,3.72,1.45,"11,658 m/s","slow-spiral Δv","Required Earth velocity change",GOLD),
    (4.80,1.86,3.72,1.45,"1.697×10³³ J","ΔE orbital","Earth–Moon orbital-energy increase",TEAL),
    (8.78,1.86,3.72,1.45,"1.664×10³⁵ J","minimum draw","Launcher energy estimate",BLUE),
    (.82,3.52,3.72,1.45,"1.415×10²² kg","reaction mass","0.193 lunar masses; 3×10²² kg available",GOLD),
    (4.80,3.52,3.72,1.45,"1.114×10²⁵ W","electricity at 2.7 AU","versus 6.338×10²⁴ W peak launcher demand",TEAL),
    (8.78,3.52,3.72,1.45,"1.90×10¹⁷ W","mirror capacity","≈1.10× today’s Earth sunlight at target",BLUE),
]
for x,y,w,h,big,small,desc,accent in metrics:
    rect(s,x,y,w,h,PANEL,LINE)
    textbox(s,x+.18,y+.12,w-.36,.45,big,21,accent,True)
    textbox(s,x+.18,y+.60,w-.36,.25,small,10,MUTED,True)
    textbox(s,x+.18,y+.92,w-.36,.44,desc,10,WHITE)
rect(s,.82,5.26,11.68,1.25,PANEL2,LINE)
textbox(s,1.02,5.40,3.2,.28,"PV FIELD · CONFIGURED",10,TEAL,True)
textbox(s,1.02,5.74,11.1,.58,"5 million panels · 0.5 AU · 30% solar interception · 2.109×10²² m² area · 64,950 km square side per panel · estimated panel mass 2.109×10²¹ kg.\nEstimated electric delivery: 1.290×10²⁵ W at 1 AU; 1.114×10²⁵ W at 2.7 AU. Deployment and safe transfer are unverified.",10,WHITE)

# 5 — ML / simulation
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
title(s, 5, "Model & ML", "Policy chooses commands; physics computes outcomes", "Two simulation speeds, separate physical metrics, and two distinct post-arrival checks.", "3:20–4:25")
# top pipeline
boxes=[(.82,"31 observations","orbit · mass · climate"),(4.79,"PPO policy","5 continuous actions"),(8.76,"physics transition","ejection · recoil · climate")]
for x,h,b in boxes:
    rect(s,x,1.91,3.62,1.05,PANEL,LINE)
    textbox(s,x+.14,2.05,3.34,.30,h,14,TEAL,True,align=PP_ALIGN.CENTER)
    textbox(s,x+.14,2.42,3.34,.28,b,10,MUTED,align=PP_ALIGN.CENTER)
line(s,4.45,2.43,4.73,2.43,GOLD,2)
line(s,8.42,2.43,8.70,2.43,GOLD,2)
# engines
card_text(s,.82,3.18,5.72,1.30,"FAST · TRAINING", "Kepler-based propagation; Moon orbit included; documented approximations. Used for PPO experience collection.", TEAL, 13, 10)
card_text(s,6.78,3.18,5.72,1.30,"HIGH-FIDELITY · N-BODY", "Sun, Earth, Moon, fragments, Mars, Mercury, Venus, Jupiter; conservation/event checks. Beyond Jupiter omitted.", BLUE, 13, 10)
card_text(s,.82,4.72,5.72,1.20,"ACTIVE HOLD · agent still controls", "Default 200 consecutive simulated years in tolerance; a breach resets progress.", GOLD, 12, 10)
card_text(s,6.78,4.72,5.72,1.20,"PASSIVE COAST · separate test", "Zero ejections in a high-fidelity coast after evaluation; not agent control or PPO feedback.", TEAL, 12, 10)
rect(s,.82,6.15,11.68,.50,"3D2531","A66C57")
textbox(s,1.02,6.26,11.20,.25,"EVIDENCE GAP: no current-version PPO retraining or target-hold evaluation. No current RL success is claimed.",10,RED,True,align=PP_ALIGN.CENTER)

# 6 — visuals / target criteria
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
title(s, 6, "Views & criteria", "A clear label is part of the model", "Separate state-driven physics from explanatory visualization and project-internal thresholds.", "4:25–5:20")
card_text(s,.82,1.88,5.62,2.16,"STATE-DRIVEN OUTPUTS", "• Orbital states and ejection events\n• Finite packet / fragment fates and recoil\n• Energy and mass ledger\n• Climate temperature / ice-line indicators\n• 12 representative collector tracks", TEAL,13,11)
card_text(s,6.66,1.88,5.84,2.16,"ILLUSTRATIVE / SIMPLIFIED", "• Selectable Earth POV atmosphere / space-weather layers\n• Clouds, winds, storms; not coupled climate\n• Overview scale enlargement / compressed distances\n• 5 million collector hardware is not tracked unit-by-unit\n• Arm meshes are not stress or survival analysis", GOLD,13,11)
rect(s,.82,4.30,11.68,1.58,PANEL2,LINE)
label(s,1.02,4.47,5.4,"Internal orbit / climate windows · not official limits",BLUE)
textbox(s,1.02,4.84,11.15,.82,"2.7 AU ±0.05 AU · orbital-plane inclination ≤5° (degrees, not axial tilt) · absolute final eccentricity ≤0.05 (not Δe) · climate target 288 K ±5 K.\nThe reviewed PDFs contain no numeric rotation-axis-tilt or eccentricity-change limits. Spin, obliquity and seasons are not modeled.",11,WHITE)
textbox(s,1.02,6.12,11.2,.48,"Mars safety shield: configurable tide / Moon-orbit / minimum-distance screens; internal settings, not official safety standards. Default safe distance ≈0.0246 AU.",10,MUTED)

# 7 — four independent failed-model screens
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
title(s, 7, "Four model screens", "Four concepts tested against the same target", "Model 1 merges laser + Dyson; model 4 is one explicit ideal thruster case. Main launcher remains separate.", "5:20–6:10")
card_text(s,.82,1.92,5.72,1.90,"FAIL 1 · LASER + DYSON", "Required beam: 6.614×10²⁵ W\nSimple global estimate: ≈43,993 K\nReport-scale beam supplies 19.5% of F_req.\nArm capacity / SF limit are illustrative proxies.", RED,12,10)
card_text(s,6.78,1.92,5.72,1.90,"FAIL 2 · HYDROGEN", "Ideal fusion screen needs 1.949×10²¹ kg.\nA 100 R_E scoop collects 1.346×10¹⁵ kg in 10,000 y.\nEven the whole solar wind supplies only 24.3%.\n1-AU flux held constant: optimistic.", GOLD,12,10)
card_text(s,.82,4.02,5.72,1.90,"FAIL 3 · LORENTZ ARMS", "F_req at 1 AU: 4.351×10¹⁷ A per arm.\nAssumed 45 T / 1 m-cable screen: 2.25×10⁸ A.\nInside Earth's magnetosphere, actual net force is zero.\nExternal-field run is hypothetical.", BLUE,12,10)
card_text(s,6.78,4.02,5.72,1.90,"FAIL 4 · EARTH-DERIVED THRUSTERS", "v_e = 10 km/s; selected Earth fuel = 1×10²³ kg.\nRocket equation: 168.9 m/s → 1.01144 AU.\nIdeal requirement: 4.111×10²⁴ kg (68.8% of Earth).\nEven the 1×10²⁴ kg slider ceiling stays short.", RED,12,10)
rect(s,.82,6.12,11.68,.57,PANEL2,LINE)
textbox(s,1.02,6.25,11.18,.30,"Four simplified, independent screens—not validated hardware models and not part of the main launcher. The main concept still has exactly three individual enablers.",10,MUTED,True,align=PP_ALIGN.CENTER)

# 8 — validation and next gate
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
title(s, 8, "Evidence & next gate", "Passing checks is not passing the mission", "The honest contribution today is inspectable bookkeeping, bounded assumptions and explicit gaps.", "6:10–7:00")
for x,big,small,acc in [(.82,"48 / 49","tests passed¹",TEAL),(4.80,"TORCH","1 device benchmark blocked",GOLD),(8.78,"BLOCKED","physics check: SB3 / Torch",RED)]:
    rect(s,x,1.90,3.72,1.08,PANEL,LINE)
    textbox(s,x+.14,2.02,3.44,.42,big,22,acc,True,align=PP_ALIGN.CENTER)
    textbox(s,x+.14,2.49,3.44,.24,small,10,MUTED,True,align=PP_ALIGN.CENTER)
textbox(s,.82,3.02,11.68,.20,"¹ 48 tests pass when the torch-dependent device benchmark is deselected; the full run stops at the missing torch import. No current physics-check run.",8,MUTED,align=PP_ALIGN.CENTER)
card_text(s,.82,3.28,5.72,2.34,"OPEN · NOT VERIFIED", "• Current PPO hold performance\n• Main-launcher structural SF ≥1.5\n• Launch shock / Earth damage\n• Radiator design / beam safety\n• Asteroid capture / deployment\n• Long-term orbit / full climate", RED,13,10)
card_text(s,6.78,3.28,5.72,2.34,"NEXT EVIDENCE GATE", "1. Retrain current version; evaluate held-out seeds through active hold.\n2. Separately evaluate passive N-body coast.\n3. Publish physical metrics apart from reward.\n4. Add structural, thermal, sensing and launch-shock analysis.", TEAL,13,10)
rect(s,.82,5.92,11.68,.72,PANEL2,LINE)
textbox(s,1.02,6.10,11.18,.38,"Guide format: demo ≤3 min · pitch 7 min · Q&A 5 min · report ≤8 pages. Structural requirement: SF ≥1.5; no main-launcher pass is claimed.",10,WHITE,True,align=PP_ALIGN.CENTER)

prs.core_properties.title = "OCTO-EARTH Pitch Deck"
prs.core_properties.subject = "Seven-minute pitch on a hypothetical planetary mobility simulation"
prs.core_properties.author = "OCTO-EARTH project"
prs.core_properties.keywords = "OCTO-EARTH, planetary propulsion, simulation, physics, reinforcement learning"
prs.save(OUT)
print(f"Wrote {OUT}")
